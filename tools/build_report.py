"""Build the editable Russian report and its printable HTML/PDF versions.

Dependencies: python-docx, Pillow, playwright (with Chromium), pypdf.
Run after tools/check_forms.py and tools/validate_html.py:
    python tools/build_report.py

Evidence is read from the existing JSON files; no test result is invented.
The screenshots are embedded unchanged. The PDF is printed from the same
report content as the DOCX, rather than requiring Microsoft Word automation.
"""

from __future__ import annotations

import argparse
import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
STEM = "Assignment4_Aitpayev_Tamerlan"
REPOSITORY = "https://github.com/tamerlanaitpayev1-netizen/web4"
SOURCE_LINKS = [
    ("Bootstrap 5.3: формы и компоненты", "https://getbootstrap.com/docs/5.3/forms/overview/"),
    ("Bootstrap 5.3: сетка", "https://getbootstrap.com/docs/5.3/layout/grid/"),
    ("MDN: HTML constraint validation", "https://developer.mozilla.org/en-US/docs/Web/HTML/Guides/Constraint_validation"),
    ("MDN: input type=time и шаг в секундах", "https://developer.mozilla.org/en-US/docs/Web/HTML/Reference/Elements/input/time"),
    ("MDN: :user-invalid", "https://developer.mozilla.org/en-US/docs/Web/CSS/Reference/Selectors/:user-invalid"),
    ("MDN: связь label с элементом формы", "https://developer.mozilla.org/en-US/docs/Web/HTML/Reference/Elements/label"),
    ("W3C Nu HTML Checker", "https://validator.w3.org/nu/"),
]


def set_cell_background(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def add_hyperlink(paragraph, title: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0F6B5F")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.6)
        section.left_margin, section.right_margin = Cm(1.6), Cm(1.6)
        section.header_distance, section.footer_distance = Cm(.65), Cm(.65)
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.font.color.rgb = RGBColor.from_string("20372F")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08
        for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 13)):
            styles[name].font.name = "Calibri"
            styles[name].font.size = Pt(size)
            styles[name].font.color.rgb = RGBColor.from_string("0F6B5F")
            styles[name].paragraph_format.space_before = Pt(7)
            styles[name].paragraph_format.space_after = Pt(7)
        styles["Caption"].font.name, styles["Caption"].font.size = "Calibri", Pt(8)
        styles["Caption"].font.color.rgb = RGBColor.from_string("65736B")
        styles["Caption"].paragraph_format.space_after = Pt(6)
        self.doc.core_properties.title = "Assignment 4 — Forms and Validation"
        self.doc.core_properties.author = "Aitpayev Tamerlan"
        self.doc.core_properties.subject = "HTML + Bootstrap, группа IT-2510"
        header = section.header.paragraphs[0]
        header.text = "AITU  /  WEB TECHNOLOGIES  /  ASSIGNMENT 4"
        header.runs[0].font.size = Pt(7)
        header.runs[0].font.color.rgb = RGBColor.from_string("65736B")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Aitpayev Tamerlan · IT-2510     |     ").font.size = Pt(8)
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), "PAGE")
        footer._p.append(page)
        self.parts = []

    def heading(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level=level)
        self.parts.append(f"<h{level + 1}>{html.escape(text)}</h{level + 1}>")

    def paragraph(self, text: str, caption: bool = False) -> None:
        self.doc.add_paragraph(text, style="Caption" if caption else None)
        self.parts.append(f'<p class="{"caption" if caption else ""}">{html.escape(text)}</p>')

    def link(self, title: str, url: str) -> None:
        add_hyperlink(self.doc.add_paragraph(), title, url)
        self.parts.append(f'<p class="source"><a href="{html.escape(url, quote=True)}">{html.escape(title)}</a></p>')

    def page_break(self) -> None:
        self.doc.add_page_break()
        self.parts.append('<div class="page-break"></div>')

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 9) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for cell, title in zip(table.rows[0].cells, headers):
            cell.text = title
            set_cell_background(cell, "0F6B5F")
        header_repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
        for cells in rows:
            row = table.add_row()
            for cell, value in zip(row.cells, cells):
                cell.text = value
        for index, row in enumerate(table.rows):
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                if index > 0 and index % 2 == 0:
                    set_cell_background(cell, "F2F6F3")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.02
                    for run in paragraph.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(font_size)
                        if index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        html_rows = ["<thead><tr>" + "".join(f"<th>{html.escape(h)}</th>" for h in headers) + "</tr></thead><tbody>"]
        for row in rows:
            html_rows.append("<tr>" + "".join(f"<td>{html.escape(c).replace(chr(10), '<br>')}</td>" for c in row) + "</tr>")
        columns = "".join(f'<col style="width:{width / sum(widths) * 100:.2f}%">' for width in widths)
        self.parts.append(f'<table style="font-size:{font_size}pt"><colgroup>{columns}</colgroup>' + "".join(html_rows) + "</tbody></table>")

    def picture(self, relative: str, caption: str, width: float = 6.7, max_height: float = 7.6) -> None:
        path = ROOT / relative
        with Image.open(path) as picture:
            image_width, image_height = picture.size
        width = min(width, max_height * image_width / image_height)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run().add_picture(str(path), width=Inches(width))
        self.paragraph(caption, caption=True)
        # Replace the just-added HTML caption with an image + its caption.
        self.parts.insert(len(self.parts) - 1, f'<img class="report-image" style="width:{width}in" src="../{html.escape(relative)}" alt="{html.escape(caption, quote=True)}">')

    def picture_pair(self, left: str, right: str, left_caption: str, right_caption: str) -> None:
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column in table.columns:
            column.width = Cm(8.9)
        html_cells = []
        for cell, relative, caption in zip(table.rows[0].cells, [left, right], [left_caption, right_caption]):
            cell.width = Cm(8.9)
            with Image.open(ROOT / relative) as picture:
                width, height = picture.size
            image_inches = min(2.55, 7.95 * width / height)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(caption).bold = True
            paragraph = cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(ROOT / relative), width=Inches(image_inches))
            html_cells.append(f'<div><p class="pair-caption">{html.escape(caption)}</p><img style="width:{image_inches}in" src="../{html.escape(relative)}" alt="{html.escape(caption, quote=True)}"></div>')
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        self.parts.append('<div class="picture-pair">' + "".join(html_cells) + "</div>")

    def save(self, pdf: bool = True) -> dict:
        docx_path = DOCS / f"{STEM}.docx"
        html_path = DOCS / "report.html"
        pdf_path = DOCS / f"{STEM}.pdf"
        self.doc.save(docx_path)
        document = '''<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>Assignment 4 — Aitpayev Tamerlan</title><style>
@page { size: A4; margin: 16mm; }
* { box-sizing: border-box; }
body { margin: 0; font: 11pt/1.12 Arial, sans-serif; color: #20372f; }
h1 { font-size: 26pt; color: #0f6b5f; margin: 0 0 10pt; }
h2 { font-size: 17pt; color: #0f6b5f; margin: 7pt 0; }
h3 { font-size: 13pt; color: #0f6b5f; margin: 7pt 0; }
p { margin: 0 0 6pt; }
a { color: #0f6b5f; overflow-wrap: anywhere; }
table { width: 100%; table-layout: fixed; border-collapse: collapse; margin: 8pt 0; line-height: 1.1; }
th, td { text-align: left; vertical-align: top; padding: 5pt; border: .5pt solid #bdcdc3; overflow-wrap: anywhere; }
th { background: #0f6b5f; color: white; }
tbody tr:nth-child(even) { background: #f2f6f3; }
thead { display: table-header-group; }
tr { break-inside: avoid; }
.page-break { break-before: page; }
.caption { font-size: 8pt; line-height: 1.2; color: #65736b; margin: 4pt 0 7pt; }
.report-image { display: block; max-width: 100%; height: auto; margin: 8pt auto 3pt; }
.picture-pair { display: flex; align-items: flex-start; gap: 14pt; justify-content: center; }
.picture-pair > div { flex: 1; text-align: center; }
.picture-pair img { max-width: 100%; height: auto; }
.pair-caption { font-size: 9pt; font-weight: bold; min-height: 24pt; }
.source { margin-bottom: 5pt; font-size: 9pt; }
@media screen { body { max-width: 210mm; padding: 16mm; margin: auto; background: #fff; } html { background: #e8efea; } .page-break { margin-top: 28pt; border-top: 1px solid #d7e1da; padding-top: 20pt; } }
</style></head><body>''' + "\n".join(self.parts) + "</body></html>"
        html_path.write_text(document, encoding="utf-8")
        evidence = {"docx": str(docx_path.relative_to(ROOT)), "html": str(html_path.relative_to(ROOT))}
        if pdf:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=True)
                page = browser.new_page()
                page.goto(html_path.as_uri(), wait_until="networkidle")
                page.emulate_media(media="print")
                page.pdf(path=str(pdf_path), format="A4", print_background=True, prefer_css_page_size=True)
                browser.close()
            evidence["pdf"] = str(pdf_path.relative_to(ROOT))
            from pypdf import PdfReader
            reader = PdfReader(pdf_path)
            evidence["pdf_pages"] = len(reader.pages)
            evidence["pdf_text_characters"] = sum(len(page.extract_text() or "") for page in reader.pages)
        return evidence


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def concise_audit(audit_text: str) -> list[list[str]]:
    # The full source audit remains linked. These rows preserve every issue
    # and its exact original line, while keeping the submitted table readable.
    expected_numbers = re.findall(r"^\| (\d+) \|", audit_text, re.MULTILINE)
    assert expected_numbers == [str(n) for n in range(1, 18)], "The Task 3 source audit changed; review report rows."
    return [
        ["1", '3: <input type="text" placeholder="Your name">', "У имени нет постоянной подписи; placeholder исчезает при вводе.", 'Добавлены id="full-name" и связанный видимый label.'],
        ["2", '4: <input type="text" placeholder="Email">', "Назначение заполненного поля email нельзя понять по постоянной подписи.", 'Добавлены id="email" и label for="email".'],
        ["3", '4: type="text" у Email', "Вместо адреса принимается произвольная строка; нет подходящей email-клавиатуры.", 'Используются type="email", autocomplete="email" и required.'],
        ["4", '3, 4, 6, 8, 10, 14, 18, 19: нет name; например <input type="text" id="phone">', "Значения не включаются в обычную отправку формы. Атрибут id не заменяет name.", "Каждому полю задано осмысленное name: full_name, email, phone, guests, date, meal, comments, terms."],
        ["5", "5: <label>Phone</label>", "Нажатие на подпись не переносит фокус; программной связи с полем нет.", 'У подписи задан for="phone", совпадающий с id поля.'],
        ["6", '6: <input type="text" id="phone">', "Браузер не знает, что поле предназначено для телефона, и может открыть обычную клавиатуру.", 'Используются type="tel" и autocomplete="tel". Сам tel не проверяет формат номера.'],
        ["7", '7–8: for="guests", но id="people"', "Подпись Guests не связана с полем количества гостей.", 'Поле переименовано в id="guests"; for и id совпадают.'],
        ["8", '8: type="text" у Guests', "Можно ввести буквы, ноль, отрицательное или дробное число гостей.", 'Заданы type="number" min="1" step="1" required.'],
        ["9", '10: type="text" id="date" placeholder="dd/mm/yyyy"', "Подсказка не проверяет существование даты; можно ввести произвольный текст.", 'Используется type="date" name="date" required с нативным вводом даты.'],
        ["10", '10 и 12: повторяется id="date"', "Повторный id нарушает уникальность и делает связи элементов неоднозначными.", 'У даты остаётся id="date"; у Inside — seating-inside; у Terrace — seating-terrace.'],
        ["11", '12–13: name="inside" и name="terrace"', "Оба варианта можно выбрать одновременно: это разные radio-группы.", 'Общее name="seating" и разные value="inside" / "terrace".'],
        ["12", "12–13: текст > Inside и > Terrace", "Нажатие на текст не выбирает вариант; у radio нет надёжной подписи.", "Каждый вариант помещён в label, дополнительно связанный через for и уникальный id."],
        ["13", "11–13: <p>Seating:</p> перед radio", "Текстовый заголовок не задаёт семантическую группу для вспомогательных технологий.", "Варианты объединены fieldset с legend Seating."],
        ["14", "14: <select>", "Не объясняется назначение списка Lunch / Dinner; нет связанной подписи.", 'Добавлены id="meal", name="meal", label и пустая начальная опция.'],
        ["15", '<textarea placeholder="Comments"></textarea> (18)', "Подсказка комментария исчезает при вводе; постоянной подписи нет.", "Добавлены id, name и связанный label. Комментарий помечен как необязательный."],
        ["16", '19: <input type="checkbox"> I agree to the terms', "Текст не переключает checkbox, отсутствует связанное доступное имя.", 'Добавлены id="terms", name, value, label for="terms" и required. Условия доступны в details.'],
        ["17", '20: <div class="btn btn-primary">Reserve</div>', "Элемент похож на кнопку, но не отправляет форму и не имеет нативного клавиатурного поведения.", 'Заменён на <button type="submit">; работают клавиатура, валидация и отправка.'],
    ]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-pdf", action="store_true", help="Create DOCX and HTML only.")
    args = parser.parse_args()
    required = [
        "docs/task2-results.json", "docs/task3-audit.md", "docs/keyboard-results.json",
        "screenshots/task1-1280.png", "screenshots/task1-375.png",
        "screenshots/task2-overview.png", "screenshots/task2-errors-detail.png",
        "screenshots/task3-1280.png", "reference/page-2-image-3-Image3.jpg",
        "docs/w3c-task1.json", "docs/w3c-task2.json", "docs/w3c-task3.json",
    ]
    missing = [relative for relative in required if not (ROOT / relative).exists()]
    if missing:
        raise SystemExit("Required evidence not ready: " + ", ".join(missing))
    results = load_json(DOCS / "task2-results.json")
    checks = results["results"]
    keyboard = load_json(DOCS / "keyboard-results.json")
    examples = [item for item in checks if item["test"].startswith("PDF example:")]
    boundaries = [item for item in checks if item["test"].startswith("Boundary:")]
    assert len(examples) == 32, "Expected all 32 PDF examples."
    audit_rows = concise_audit((DOCS / "task3-audit.md").read_text(encoding="utf-8"))
    report = Report()

    report.doc.add_heading("Assignment 4", 0)
    report.parts.append("<h1>Assignment 4</h1>")
    report.heading("Формы и валидация · HTML + Bootstrap", 2)
    report.paragraph("Студент: Aitpayev Tamerlan  ·  Группа: IT-2510  ·  Дата: 17 октября 2026")
    report.link("GitHub: " + REPOSITORY, REPOSITORY)
    report.paragraph("В работе реализованы три формы: покупка билета по образцу, регистрация на хакатон с нативными ограничениями и исправленная форма кафе. Bootstrap и Montserrat сохранены локально. В страницах нет JavaScript и атрибута novalidate.")
    report.heading("1. Бронирование железнодорожного билета")
    report.paragraph("Макет построен на row, col-12, col-md-6 и col-md-4. Использованы form-control, form-select, btn-check, form-check, form-switch и input-group. Собственный CSS задаёт цвета из задания, скругления, тень и отступы; фиксированных ширин и абсолютного позиционирования в CSS этого задания нет.")
    report.picture("screenshots/task1-1280.png", "Рисунок 1. Task 1 — ширина окна 1280 px. Показан пример с выбранными One way, Coupé, 1 adult и введёнными Astana / Almaty; это не исходное пустое состояние.", width=6.75, max_height=5.6)
    report.paragraph("Обязательны маршрут, дата отправления, класс вагона, телефон, email и согласие; список пассажиров содержит 1–6 взрослых. Return необязателен. Подписи связаны с полями через for/id. Нативная кнопка reset очищает введённые значения и выбор, поскольку исходная форма не содержит value/checked по умолчанию; служебный префикс телефона +7 сохраняется.")

    report.page_break()
    report.heading("Task 1 — сравнение на телефоне, 375 px")
    report.paragraph("На узком экране колонки становятся вертикальными. Классы flex-column-reverse flex-md-row размещают Find trains выше Clear на телефоне, сохраняя порядок элементов в HTML. Ниже исходный мобильный образец и результат при ширине окна 375 px.")
    report.picture_pair("reference/page-2-image-3-Image3.jpg", "screenshots/task1-375.png", "Образец из Assignment 4.pdf\nМакет для 375 px", "Результат task1.html\nОкно 375 px")

    report.page_break()
    report.heading("2. Регистрация на AITU Hack 2026")
    report.paragraph("Форма использует подходящие type и атрибуты required, min, max, step, minlength, maxlength. Правила не зависят от JavaScript; pattern не используется. Все поля обязательны, кроме GitHub. Email проверяется как адрес без ограничения домена. step=1800 у времени означает 1800 секунд, то есть 30 минут.")
    report.picture("screenshots/task2-overview.png", "Рисунок 2. Task 2 — окно 1280 × 1000 px, верхняя часть исходной формы. При открытии сообщения об ошибках скрыты.", width=6.65, max_height=5.0)
    report.picture("screenshots/task2-errors-detail.png", "Рисунок 3. Ошибочные имя A, email a.fedenko@ и возраст 15 после ввода и смены фокуса: красная граница и текст .invalid-feedback.", width=5.7, max_height=2.6)
    report.paragraph("CSS показывает сообщения через :user-invalid после взаимодействия с полем. Для составных групп используется :has(). Для клавиатуры предусмотрен видимый фокус; пояснения и сообщения связаны через aria-describedby.")

    example_rows = []
    labels = {
        "Full name": "Имя", "Email": "Email", "Age": "Возраст", "Team size": "Команда",
        "Track": "Трек", "Arrival": "Прибытие", "Check-in": "Время", "Entry fee": "Взнос",
        "Experience": "Опыт", "GitHub": "GitHub", "Password": "Пароль",
        "Motivation": "Мотивация", "Agreement": "Согласие",
    }
    example_values = {
        "Choose...": "Пустая опция Choose…", "none": "Ничего не выбрано", "empty": "Пустое поле",
        "full HTTPS URL": "https://github.com/yerassyl-dev", "no protocol": "github.com/yerassyl-dev",
        "120 characters": "120 символов", "49 characters": "49 символов", "checked": "Установлено", "unchecked": "Не установлено",
    }
    for item in examples:
        title = item["test"].removeprefix("PDF example: ")
        field, value = title.split(": ", 1)
        detail = item["detail"]
        expected = "Принять" if detail["expected_valid"] else "Отклонить"
        actual = "Принято" if detail["actual_valid"] else "Отклонено"
        actual += f"; submit={detail['submit_events']}"
        actual += " ✓" if item["passed"] else " — НЕ СОВПАЛО"
        example_rows.append([labels.get(field, field), example_values.get(value, value), expected, actual])
    report.page_break()
    report.heading("Task 2 — результаты примеров из задания (1/2)")
    report.paragraph(f"Браузер: {results['browser']}. Все примеры проверены вводом значений и нажатием submit. Для каждого примера остальные обязательные поля заполнены корректно. Итог примеров PDF: {sum(item['passed'] for item in examples)} из {len(examples)} совпали с ожидаемым результатом.")
    report.table(["Поле", "Проверяемое значение", "Ожидалось", "Фактически"], example_rows[:18], [2.5, 7.0, 2.6, 5.7], 9.5)
    report.paragraph("submit=0 означает, что браузер заблокировал отправку и событие submit не возникло; submit=1 означает, что событие возникло у валидной формы. Во время этих проверок тестовый обработчик отменял фактическую отправку. В сдаваемом HTML этого обработчика нет.")
    report.heading("Как заданы ограничения", 2)
    report.paragraph("Длина имени 2–40, пароля 8–20, мотивации 50–300 задана minlength/maxlength. Числа ограничены min/max: возраст 16–30, команда 1–4, взнос 0–5000. step=1 разрешает только целые числа; step=500 задаёт шаг взноса. Даты ограничены 2026-11-20…2026-11-22, время — 09:00…18:00. У select пустая опция имеет value=\"\"; radio опыта имеют одно name. Обязательное согласие реализовано checkbox с required.")

    report.page_break()
    report.heading("Task 2 — результаты примеров из задания (2/2)")
    report.table(["Поле", "Проверяемое значение", "Ожидалось", "Фактически"], example_rows[18:], [2.5, 7.0, 2.6, 5.7], 9.5)
    report.heading("Границы и поведение формы", 2)
    report.paragraph(f"Дополнительно выполнены {len(boundaries)} проверки граничных и новых значений; успешно: {sum(item['passed'] for item in boundaries)}. Проверены края диапазонов (16/30, 1, 20/22 ноября, 09:00/18:00, 0/5000), значения вне диапазонов, дробный возраст, длины 2/40, 8/20, 50/300 и все четыре трека.")
    max_checks = [item for item in checks if item["test"].startswith("maxlength prevents")]
    report.paragraph(f"Проверки ограничения ввода maxlength: {sum(item['passed'] for item in max_checks)}/{len(max_checks)}. При попытке набрать лишний символ браузер оставляет длину 40, 20 или 300. Это проверяется реальным вводом, а не только присваиванием значения через программный интерфейс.")
    report.paragraph(f"Общий браузерный набор: {results['passed']}/{results['total']} проверок. Он включает HTML-атрибуты, примеры, границы, отсутствие ошибок при открытии, появление ошибки после изменения, запрет script/novalidate и отсутствие горизонтального переполнения при 375 и 1280 px. Отдельная проверка клавиатуры и семантики трёх форм: {keyboard['passed']}/{keyboard['total']}. Полные данные: docs/task2-results.json и docs/keyboard-results.json; сценарии находятся в tools/.")
    report.paragraph("Клиентская проверка подтверждает формат и указанные ограничения, но не существование email или страницы GitHub. Реальная регистрация, аккаунт и платёж не создаются.")

    report.page_break()
    report.heading("3. Исправление формы бронирования кафе")
    report.paragraph("В task3.html сохранены все поля исходника: имя, email, телефон, гости, дата, Inside / Terrace, Lunch / Dinner, комментарий, согласие и Reserve. Добавлены видимые связанные подписи, name для отправки значений, уникальные id, подходящие типы, группы fieldset/legend и нативная submit-кнопка.")
    report.picture("screenshots/task3-1280.png", "Рисунок 4. Исправленная форма кафе — ширина окна 1280 px. Мобильный снимок 375 px также сохранён в screenshots/task3-375.png.", width=6.6, max_height=7.0)
    report.paragraph("Телефон и комментарий необязательны; количество гостей — положительное целое число. Верхняя граница вместимости и ограничения по дате не выдуманы: в задании их нет. Для клавиатуры доступны Tab, выбор radio стрелками, checkbox пробелом и нативная кнопка отправки.")

    report.page_break()
    report.heading("Task 3 — таблица проблем исходного кода (1/2)")
    report.paragraph("Нумерация относится к 21 строке исходного фрагмента из PDF, переписанного в reference/task3-original.txt. Полный развёрнутый аудит находится в docs/task3-audit.md. Указаны ошибки разметки, доступности и поведения; не все из них обнаруживает HTML-валидатор.")
    report.table(["№", "Строка / точный фрагмент", "Проблема для пользователя", "Исправление"], audit_rows[:9], [1.0, 5.0, 5.5, 6.3], 9)

    report.page_break()
    report.heading("Task 3 — таблица проблем исходного кода (2/2)")
    report.table(["№", "Строка / точный фрагмент", "Проблема для пользователя", "Исправление"], audit_rows[9:], [1.0, 5.0, 5.5, 6.3], 9)
    report.paragraph("Отсутствие doctype, head и body не названо ошибкой исходника: в PDF показан фрагмент. В итоговой самостоятельной странице эти элементы, lang=\"en\", UTF-8 и viewport добавлены. Отсутствие action/method у исходного form также допустимо: браузер использует текущий адрес и GET.")

    report.page_break()
    report.heading("Task 3 — альтернативы и подтверждение проверки")
    report.heading("Два альтернативных исправления", 2)
    report.paragraph("1. Guests, строки 7–8. Можно оставить id=\"people\" и заменить подпись на for=\"people\" — это минимальное корректное изменение. Выбрано id=\"guests\", чтобы for, id и name назывались одинаково. В автономном учебном фрагменте нет зависимостей от старого id; в существующем проекте перед переименованием потребовалось бы проверить такие зависимости.")
    report.paragraph("2. Reserve, строка 20. Вместо button допустим <input type=\"submit\" value=\"Reserve\">. Выбран button type=\"submit\", потому что он позволяет поместить внутри текст и декоративную стрелку с aria-hidden, сохраняя нативную отправку и доступное имя. Одни role=\"button\" и tabindex на div не добавили бы отправку и клавиатурное поведение.")
    validation_rows = []
    for task in ("task1", "task2", "task3", "index", "completed"):
        evidence = load_json(DOCS / f"w3c-{task}.json")
        metadata = evidence["_evidence"]
        status = metadata.get("status", "unknown")
        messages = evidence.get("messages", [])
        warnings = sum(message.get("type") == "info" and message.get("subType") == "warning" for message in messages)
        if status in ("passed", "failed"):
            result = f"Ошибок: {metadata['error_count']}; предупреждений: {warnings}."
            result += " Проверка пройдена." if status == "passed" else " Требуются исправления."
        else:
            result = "Проверка недоступна: " + metadata.get("reason", status) + ". Успех не заявляется."
        validation_rows.append([task + ".html", result, f"docs/w3c-{task}.json"])
    report.heading("W3C Nu HTML Checker", 2)
    report.table(["Страница", "Фактический результат", "Сохранённое доказательство"], validation_rows, [3, 8.5, 6.3], 9)
    report.paragraph("В отчёте использованы сохранённые ответы официального валидатора, а не предположение о корректности. Поле _evidence в JSON содержит URL валидатора, время проверки и количество ошибок. Проверка HTML дополняет проверку клавиатурой и поведения формы.")
    report.heading("Запуск и пределы демонстрации", 2)
    report.paragraph("Запуск: python serve.py, затем http://127.0.0.1:8000/. Формы отправляют POST на completed.html; локальный сервер считывает и отбрасывает введённые данные, затем перенаправляет на страницу завершения. Полноценного backend, хранения, выдачи билетов или бронирования нет. Для реального сервиса понадобились бы серверная валидация и бизнес-логика.")
    report.heading("Вывод", 2)
    report.paragraph("Работа показывает, как собрать адаптивные формы Bootstrap, выразить ограничения нативными HTML-атрибутами и объяснить ошибку после взаимодействия с полем. Связанные подписи, общие radio-группы, name и настоящие кнопки одновременно улучшают доступность и правильность отправки. Проверка границ дополняет примеры из задания и помогает подготовиться к новым значениям на защите.")
    report.heading("Источники", 2)
    report.paragraph("Исходное задание и изображения: Assignment 4.pdf, предоставленный для этой работы. Технические источники:")
    for title, url in SOURCE_LINKS:
        report.link(title, url)
    output = report.save(pdf=not args.no_pdf)
    output.update({"student": "Aitpayev Tamerlan", "group": "IT-2510", "repository": REPOSITORY, "browser_checks_passed": results["passed"], "browser_checks_total": results["total"], "keyboard_checks_passed": keyboard["passed"], "keyboard_checks_total": keyboard["total"], "pdf_examples": len(examples), "audit_problems": len(audit_rows)})
    (DOCS / "report-build.json").write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(output, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
